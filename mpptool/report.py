"""Berichtserzeugung im PJM-Review-Layout – PPTX (python-pptx) und PDF (matplotlib).

Ablauf: ``build_review(snapshots)`` berechnet Kennzahlen und liefert ein
``Review``-Objekt. Daraus rendern ``export_pptx`` und ``export_pdf`` die Folien;
die GUI-Vorschau nutzt dieselben Chart-Funktionen.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from pathlib import Path

from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.figure import Figure
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Emu, Inches, Pt

from . import charts, config
from .metrics import (OverdueTrendRow, SnapshotMetrics, TimelinePoint, compute_metrics,
                      overdue_tasks, overdue_trend, task_changes, timeline)
from .model import Snapshot, TaskRecord

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)   # 16:9
FIG_SIZE = (12.5, 5.6)                             # Zoll, passend zur Folie


@dataclass
class Review:
    project_name: str
    metrics: list[SnapshotMetrics]           # chronologisch
    timeline: list[TimelinePoint]
    trend_rows: list[OverdueTrendRow]
    overdue_list: list[TaskRecord]
    changes: dict[str, list[TaskRecord]] = field(default_factory=dict)

    @property
    def current(self) -> SnapshotMetrics:
        return self.metrics[-1]

    @property
    def previous(self) -> SnapshotMetrics | None:
        return self.metrics[-2] if len(self.metrics) > 1 else None

    @property
    def prev_label(self) -> str:
        return self.previous.snapshot.display_name.split()[0] if self.previous else "-"

    @property
    def cur_label(self) -> str:
        return self.current.snapshot.display_name.split()[0]


def build_review(snapshots: list[Snapshot], project_name: str | None = None) -> Review:
    if not snapshots:
        raise ValueError("Mindestens ein Snapshot erforderlich")
    snaps = sorted(snapshots, key=lambda s: (s.snapshot_date, s.snapshot_id or 0))
    ms = [compute_metrics(s) for s in snaps]
    prev = ms[-2] if len(ms) > 1 else None
    name = project_name or Path(snaps[-1].source_file).stem
    return Review(
        project_name=name, metrics=ms, timeline=timeline(ms),
        trend_rows=overdue_trend(prev, ms[-1]), overdue_list=overdue_tasks(ms[-1]),
        changes=task_changes(prev, ms[-1]),
    )


# --------------------------------------------------------------------------
# Folien als Figuren (gemeinsame Basis für Vorschau und PDF)
# --------------------------------------------------------------------------
def figure_timeline(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    return charts.draw_timeline(fig, review.timeline, f"Master Timeline – {review.project_name}")


def figure_department_status(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    return charts.draw_department_status(fig, review.current)


def figure_overdue_by_department(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    return charts.draw_overdue_by_department(fig, review.current, review.previous)


def figure_trend_table(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    return charts.draw_overdue_trend_table(fig, review.trend_rows, review.prev_label, review.cur_label,
                                           f"Top-{config.TOP_N_OVERDUES} Overdues – {review.prev_label} vs. {review.cur_label}")


def _overdue_rows(review: Review, limit: int = 18) -> list[list[str]]:
    rows = []
    for t in review.overdue_list[:limit]:
        days = (review.current.as_of - t.finish).days if t.finish else 0
        rows.append([t.department, t.name, t.finish.isoformat() if t.finish else "-", f"{days} d",
                     f"{t.percent_complete:.0f} %"])
    return rows


def figure_overdue_list(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    fig.clear()
    ax = fig.add_subplot(111)
    ax.axis("off")
    ax.set_title(f"Überfällige Aufgaben – {review.cur_label} ({len(review.overdue_list)})", fontweight="bold", loc="left")
    rows = _overdue_rows(review)
    if not rows:
        ax.text(0.5, 0.5, "Keine überfälligen Aufgaben", ha="center", va="center")
        return fig
    tbl = ax.table(cellText=rows, colLabels=["Abteilung", "Aufgabe", "Geplantes Ende", "Verzug", "Fortschritt"],
                   loc="upper center", cellLoc="left", colWidths=[0.12, 0.5, 0.14, 0.1, 0.12])
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(8)
    tbl.scale(1, 1.25)
    for (r, c), cell in tbl.get_celld().items():
        if r == 0:
            cell.set_facecolor("#dddddd")
            cell.set_text_props(fontweight="bold")
    return fig


def figure_kpis(review: Review, fig: Figure | None = None) -> Figure:
    fig = fig or Figure(figsize=FIG_SIZE)
    fig.clear()
    m = review.current
    ax = fig.add_subplot(111)
    ax.axis("off")
    ax.set_title(f"{config.REPORT_TITLE} – {review.project_name}", fontweight="bold", loc="left", fontsize=16)
    kpis = [
        ("Stichtag", m.as_of.isoformat()),
        ("Überwachte Aufgaben", str(m.monitored)),
        ("Geplant geschlossen", str(m.planned_closed)),
        ("Tatsächlich geschlossen", str(m.actually_closed)),
        ("Erfüllungsgrad", f"{m.fulfilment_pct:.1f} %"),
        ("Überfällig (offen)", str(m.open_overdue)),
        ("Verspätet geschlossen", str(m.closed_overdue)),
        ("Hold", str(m.hold)),
        ("Fachabteilungen", str(len(m.departments))),
        ("Snapshots im Vergleich", str(len(review.metrics))),
    ]
    for i, (k, v) in enumerate(kpis):
        col, row = divmod(i, 5)
        x = 0.05 + col * 0.5
        y = 0.8 - row * 0.15
        ax.text(x, y, k, fontsize=11, color="#555555", transform=ax.transAxes)
        ax.text(x + 0.3, y, v, fontsize=14, fontweight="bold", transform=ax.transAxes)
    return fig


SLIDE_BUILDERS = [
    ("Übersicht", figure_kpis),
    ("Master Timeline", figure_timeline),
    ("Aufgaben je Fachabteilung", figure_department_status),
    ("Überfällige je Fachabteilung", figure_overdue_by_department),
    ("Top-Overdues Trend", figure_trend_table),
    ("Überfällige Aufgaben", figure_overdue_list),
]


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------
def export_pdf(review: Review, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with PdfPages(str(path)) as pdf:
        for _title, builder in SLIDE_BUILDERS:
            fig = builder(review, Figure(figsize=(13.333, 7.5)))
            _footer_fig(fig, review)
            pdf.savefig(fig)
        d = pdf.infodict()
        d["Title"] = f"{config.REPORT_TITLE} – {review.project_name}"
        d["Author"] = config.REPORT_AUTHOR
    return path


def _footer_fig(fig: Figure, review: Review) -> None:
    fig.text(0.01, 0.01, f"{config.REPORT_AUTHOR} · {config.REPORT_TITLE} · Stand {review.current.as_of.isoformat()}",
             fontsize=7, color="#888888")


# --------------------------------------------------------------------------
# PPTX
# --------------------------------------------------------------------------
def export_pptx(review: Review, path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H
    blank = prs.slide_layouts[6]

    # Titelfolie
    s = prs.slides.add_slide(blank)
    _textbox(s, config.REPORT_TITLE, Inches(0.6), Inches(2.4), Inches(12), Inches(1.0), 40, bold=True)
    _textbox(s, review.project_name, Inches(0.6), Inches(3.4), Inches(12), Inches(0.8), 24)
    _textbox(s, f"Stand {review.cur_label} · {review.current.as_of.isoformat()} · Vergleich zu {review.prev_label}",
             Inches(0.6), Inches(4.3), Inches(12), Inches(0.6), 16, color="666666")

    # KPI-Folie als native Tabelle
    s = prs.slides.add_slide(blank)
    _title(s, f"Übersicht – {review.cur_label}")
    m = review.current
    kpi = [("Stichtag", m.as_of.isoformat()), ("Überwachte Aufgaben", m.monitored),
           ("Geplant geschlossen", m.planned_closed), ("Tatsächlich geschlossen", m.actually_closed),
           ("Erfüllungsgrad", f"{m.fulfilment_pct:.1f} %"), ("Überfällig (offen)", m.open_overdue),
           ("Verspätet geschlossen", m.closed_overdue), ("Hold", m.hold)]
    _table(s, [["Kennzahl", "Wert"]] + [[k, str(v)] for k, v in kpi], Inches(0.6), Inches(1.3), Inches(6), [4.0, 2.0])

    # Chart-Folien
    for title, builder in [("Master Timeline", figure_timeline),
                           ("Aufgaben je Fachabteilung", figure_department_status),
                           ("Überfällige Aufgaben je Fachabteilung", figure_overdue_by_department)]:
        s = prs.slides.add_slide(blank)
        _title(s, title)
        png = charts.figure_to_png(builder(review, Figure(figsize=FIG_SIZE)))
        _picture(s, png, Inches(0.4), Inches(1.2), Inches(12.5))

    # Trend-Tabelle
    s = prs.slides.add_slide(blank)
    _title(s, f"Top-{config.TOP_N_OVERDUES} Overdues – {review.prev_label} vs. {review.cur_label}")
    rows = [["Fachabteilung", review.prev_label, review.cur_label, "Trend"]]
    rows += [[r.department, str(r.previous), str(r.current), f"{r.arrow} {r.delta:+d}"] for r in review.trend_rows]
    if len(rows) == 1:
        rows.append(["Keine überfälligen Aufgaben", "", "", ""])
    _table(s, rows, Inches(0.6), Inches(1.3), Inches(8), [3.2, 1.6, 1.6, 1.6], font=14)

    # Liste überfälliger Aufgaben
    s = prs.slides.add_slide(blank)
    _title(s, f"Überfällige Aufgaben – {review.cur_label} ({len(review.overdue_list)})")
    rows = [["Abteilung", "Aufgabe", "Geplantes Ende", "Verzug", "Fortschritt"]] + _overdue_rows(review, 16)
    if len(rows) == 1:
        rows.append(["Keine überfälligen Aufgaben", "", "", "", ""])
    _table(s, rows, Inches(0.5), Inches(1.2), Inches(12.3), [1.6, 6.5, 1.8, 1.2, 1.2], font=10)

    # Veränderungen seit Vorwoche
    if review.previous:
        s = prs.slides.add_slide(blank)
        _title(s, f"Veränderungen {review.prev_label} → {review.cur_label}")
        ch = review.changes
        lines = [f"Neu überfällig: {len(ch['newly_overdue'])}"] + [f"   • {t.department} – {t.name}" for t in ch["newly_overdue"][:8]]
        lines += [f"Neu geschlossen: {len(ch['newly_closed'])}"] + [f"   • {t.department} – {t.name}" for t in ch["newly_closed"][:8]]
        lines += [f"Neu hinzugekommen: {len(ch['added'])}"] + [f"   • {t.department} – {t.name}" for t in ch["added"][:6]]
        _textbox(s, "\n".join(lines), Inches(0.6), Inches(1.3), Inches(12), Inches(5.8), 12)

    for slide in prs.slides:
        _textbox(slide, f"{config.REPORT_AUTHOR} · {config.REPORT_TITLE} · Stand {review.current.as_of.isoformat()}",
                 Inches(0.4), SLIDE_H - Inches(0.45), Inches(10), Inches(0.35), 9, color="888888")
    prs.save(str(path))
    return path


def _title(slide, text: str) -> None:
    _textbox(slide, text, Inches(0.5), Inches(0.35), Inches(12.3), Inches(0.8), 26, bold=True)
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.1), Inches(12.3), Emu(19050))
    line.fill.solid()
    line.fill.fore_color.rgb = RGBColor(0x1F, 0x77, 0xB4)
    line.line.fill.background()


def _textbox(slide, text: str, x, y, w, h, size: int, bold: bool = False, color: str = "222222"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
    return tb


def _picture(slide, png: bytes, x, y, w):
    import io
    slide.shapes.add_picture(io.BytesIO(png), x, y, width=w)


def _table(slide, rows: list[list[str]], x, y, w, col_inches: list[float], font: int = 14):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), x, y, w, Inches(0.4) * len(rows))
    tbl = shape.table
    for i, ci in enumerate(col_inches):
        tbl.columns[i].width = Inches(ci)
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(font)
                p.font.bold = r == 0
                if r > 0 and c == len(row) - 1 and str(val).startswith("▲"):
                    p.font.color.rgb = RGBColor(0xD6, 0x27, 0x28)
                elif r > 0 and c == len(row) - 1 and str(val).startswith("▼"):
                    p.font.color.rgb = RGBColor(0x2C, 0xA0, 0x2C)
    return shape
